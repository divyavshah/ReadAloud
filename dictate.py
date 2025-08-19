import time
from tkinter import *
from tkinter import ttk, filedialog, messagebox
import pyttsx3
import threading
from docx import Document
import PyPDF2

# Module for dictation tool
def dictation_tool():
    stop_flag = False  # Flag to stop dictation

    def open_file():
        file_path = filedialog.askopenfilename(filetypes=[("Word Files", "*.docx"), ("PDF Files", "*.pdf")])
        if file_path:
            file_label.config(text=f"File: {file_path.split('/')[-1]}")
            try:
                if file_path.endswith(".docx"):
                    doc = Document(file_path)
                    content = "\n".join([para.text for para in doc.paragraphs])
                elif file_path.endswith(".pdf"):
                    with open(file_path, "rb") as f:
                        reader = PyPDF2.PdfReader(f)
                        content = "\n".join([page.extract_text() for page in reader.pages])
                else:
                    messagebox.showerror("Error", "Unsupported file format.")
                    return
                text_area.delete("1.0", END)
                text_area.insert("1.0", content)
            except Exception as e:
                messagebox.showerror("Error", str(e))

    def start_dictation():
        nonlocal stop_flag
        stop_flag = False  # Reset the stop flag

        def dictation_thread():
            try:
                speed_level = speed_scale.get()
                rate_mapping = {1: 100, 2: 110, 3: 120, 4: 130, 5: 150, 6: 170, 7: 180, 8: 190, 9: 200, 10: 220}
                rate = rate_mapping.get(speed_level, 150)

                text = text_area.get("1.0", END).strip().split()
                selected_text = text_area.get("sel.first", "sel.last").strip() if text_area.tag_ranges("sel") else ""

                if selected_text:
                    text = selected_text.split()

                punctuation_map = {
                    ',': 'comma', '.': 'fullstop', '?': 'question mark', '!': 'exclamation mark',
                    '(': 'openparenthesis', ')': 'close parenthesis', ':': 'colon', ';': 'semicolon',
                    '-': 'dash', '"': 'quote', "'": 'apostrophe'
                }

                text_with_punctuation = " ".join(text)
                for symbol, spoken in punctuation_map.items():
                    text_with_punctuation = text_with_punctuation.replace(symbol, f" {spoken} ")

                words = text_with_punctuation.split()
                engine = pyttsx3.init()
                engine.setProperty('rate', rate)

                for word in words:
                    if stop_flag:
                        break
                    engine.say(word)
                    engine.runAndWait()

            except Exception as e:
                messagebox.showerror("Error", str(e))

        thread = threading.Thread(target=dictation_thread)
        thread.start()

    def stop_dictation():
        nonlocal stop_flag
        stop_flag = True

    root = Tk()
    root.title("Dictation Tool")

    Label(root, text="Upload a Word or PDF file for dictation:", font=("Arial", 14)).pack(pady=10)
    file_label = Label(root, text="No file selected", font=("Arial", 12), fg="gray")
    file_label.pack(pady=5)

    Button(root, text="Upload File", command=open_file, font=("Arial", 12)).pack(pady=10)

    Label(root, text="Select dictation speed:", font=("Arial", 12)).pack(pady=10)
    speed_scale = Scale(root, from_=1, to=10, orient=HORIZONTAL, font=("Arial", 12))
    speed_scale.set(5)
    speed_scale.pack(pady=5)

    Button(root, text="Start Dictation", command=start_dictation, font=("Arial", 12)).pack(pady=10)
    Button(root, text="Stop Dictation", command=stop_dictation, font=("Arial", 12)).pack(pady=10)

    instruction_label = Label(root, text="Before starting dictation, select text if you want to hear a specific part.", font=("Arial", 10), fg="gray")
    instruction_label.pack(pady=5)

    text_frame = Frame(root)
    text_frame.pack(pady=10, fill=BOTH, expand=True)

    text_area = Text(text_frame, height=10, width=80, font=("Arial", 12), bg="lightyellow")
    text_area.pack(side=LEFT, fill=BOTH, expand=True)

    scrollbar = Scrollbar(text_frame, orient=VERTICAL, command=text_area.yview)
    scrollbar.pack(side=RIGHT, fill=Y)
    text_area.config(yscrollcommand=scrollbar.set)

    root.mainloop()

# Simplified welcome screen
def welcome_screen():
    def open_dictation_tool():
        root.destroy()
        dictation_tool()

    root = Tk()
    root.title("Welcome to Dictation Tool")

    Label(root, text="Welcome! Click below to start dictation:", font=("Arial", 16, "bold")).pack(pady=20)
    Button(root, text="Start Dictation Tool", command=open_dictation_tool, font=("Arial", 14), width=25).pack(pady=20)

    root.mainloop()

# Run the welcome screen
welcome_screen()
